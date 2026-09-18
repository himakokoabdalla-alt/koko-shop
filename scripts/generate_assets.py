#!/usr/bin/env python3
"""Generate the KOKO SHOP export bundle.

This script creates the requested archive and spreadsheet-style files that are
not practical to create directly via the GitHub file API as binary assets.
It also creates a simple logo placeholder.

Prerequisites:
  pip install openpyxl pillow

Usage:
  python scripts/generate_assets.py
"""

from __future__ import annotations

import json
import os
import shutil
import zipfile
from pathlib import Path

try:
    from openpyxl import Workbook
except ImportError as exc:  # pragma: no cover
    raise SystemExit("Please install openpyxl: pip install openpyxl") from exc

try:
    from PIL import Image, ImageDraw, ImageFont
except ImportError as exc:  # pragma: no cover
    raise SystemExit("Please install pillow: pip install pillow") from exc

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "dist"
OUT_DIR.mkdir(exist_ok=True)


def write_logo_jpg() -> Path:
    path = ROOT / "assets" / "koko-shop-logo.jpg"
    path.parent.mkdir(exist_ok=True)

    width, height = 1200, 1200
    img = Image.new("RGB", (width, height), color=(241, 248, 255))
    draw = ImageDraw.Draw(img)

    # Rounded blue gradient-like background block
    draw.rounded_rectangle((120, 120, 1080, 1080), radius=220, fill=(15, 94, 196))
    draw.rounded_rectangle((210, 210, 990, 990), radius=180, fill=(0, 185, 214))

    # K monogram
    try:
        font_large = ImageFont.truetype("DejaVuSans-Bold.ttf", 450)
        font_small = ImageFont.truetype("DejaVuSans-Bold.ttf", 110)
    except Exception:
        font_large = ImageFont.load_default()
        font_small = ImageFont.load_default()

    draw.text((500, 420), "K", anchor="mm", font=font_large, fill=(255, 255, 255))
    draw.text((600, 790), "KOKO SHOP", anchor="mm", font=font_small, fill=(255, 255, 255))
    img.save(path)
    return path


def write_excel() -> Path:
    wb = Workbook()

    ws = wb.active
    ws.title = "منتجات"
    ws.append([
        "id", "title", "description", "image", "price", "oldPrice",
        "platform", "category", "subcategory", "url", "published",
        "delivery", "cashOnDelivery", "returnPolicy", "discount",
        "badge", "buttonText", "affiliateUrl"
    ])
    ws.append([
        1, "ماوس لاسلكي احترافي", "ماوس مريح وسريع للاستخدام اليومي", "https://example.com/image1.jpg", 450, 620,
        "Amazon", "الحاسب والمكتب", "ماوس وكيبورد", "#", "TRUE", "توصيل حسب المنطقة", "TRUE", "إرجاع خلال 7 أيام", "27%", "الأكثر طلبًا", "اشترِ الآن", "#"
    ])
    ws.append([
        2, "حقيبة ظهر للرحلات", "حقيبة عملية ومريحة للرحلات", "https://example.com/image2.jpg", 760, 980,
        "Noon", "الرياضة والمغامرات", "Camping", "#", "TRUE", "توصيل سريع", "TRUE", "إرجاع خلال 10 أيام", "22%", "مميز", "اشترِ الآن", "#"
    ])

    ws = wb.create_sheet("مقالات")
    ws.append(["id", "title", "text", "body", "published"])
    ws.append([1, "كيف تختار التابلت المناسب للدراسة؟", "ابدأ بتحديد احتياجك من الأداء...", "إذا كنت تبحث عن جهاز للدراسة، فيجب أن تركز على شاشة مناسبة...", "TRUE"])
    ws.append([2, "كيف تختار ماوس لاسلكي مناسب؟", "اختر ماوسًا مريحًا...", "تأكد من نوع الاتصال، وزن الجهاز، وراحة اليد...", "TRUE"])

    ws = wb.create_sheet("إعدادات")
    ws.append(["key", "value"])
    for row in [
        ["hero_title", "KOKO SHOP"],
        ["hero_text", "اكتشف منتجات عملية، عصرية، ومفيدة للمنزل، الدراسة، الرياضة، السفر، والمطبخ."],
        ["music_url", ""],
        ["music_name", "KOKO SHOP Audio"],
        ["show_platforms", "TRUE"],
        ["amazon_disclosure", "بصفتي مشاركًا لأمازون، فإنني أكسب من عمليات الشراء المؤهلة."],
        ["payment_title", "طرق الدفع والتوصيل تختلف حسب المتجر والمنصة."],
        ["payment_methods", "الدفع عند الاستلام, بطاقات بنكية, محافظ إلكترونية, دفع آمن"],
        ["delivery_text", "التوصيل يختلف حسب المنطقة والمنصة."],
        ["site_title", "KOKO SHOP"]
    ]:
        ws.append(row)

    ws = wb.create_sheet("تصنيفات")
    ws.append(["id", "name", "icon"])
    for row in [
        [1, "الإلكترونيات", "📱"],
        [2, "الكمبيوتر والمكتب", "💻"],
        [3, "المدرسة والدراسة", "📚"],
        [4, "المنزل", "🏠"],
        [5, "الفود والمطبخ", "🍳"]
    ]:
        ws.append(row)

    ws = wb.create_sheet("رياضات")
    ws.append(["id", "name", "icon", "subcategory"])
    for row in [
        [1, "Hiking", "🥾", "حقائب, أحذية, ماء"],
        [2, "Mountain Bike", "🚴", "خوذة, قفازات, إكسسوارات"],
        [3, "Diving", "🤿", "معدات سباحة, ملابس, أحزمة"],
        [4, "Gym & Fitness", "🏋️", "أدوات, ماء, ملابس"]
    ]:
        ws.append(row)

    ws = wb.create_sheet("تواصل")
    ws.append(["name", "label", "url"])
    for row in [
        ["facebook", "Facebook", "https://www.facebook.com/profile.php?id=61593947162280"],
        ["instagram", "Instagram", "https://www.instagram.com/himakokoabdalla/"],
        ["youtube", "YouTube", "https://www.youtube.com/@kokoabdallahima"],
        ["tiktok", "TikTok", "https://www.tiktok.com/@kokoabdallahima"],
        ["threads", "Threads", "https://www.threads.com/@kokoshoppingstore"],
        ["snapchat", "Snapchat", "https://www.snapchat.com/add/koko-shoppingst"]
    ]:
        ws.append(row)

    ws = wb.create_sheet("طرق الدفع")
    ws.append(["id", "title", "icon", "description", "showOnSite"])
    for row in [
        [1, "الدفع عند الاستلام", "💵", "متاح في بعض المناطق", "TRUE"],
        [2, "بطاقات بنكية", "💳", "بطاقات Visa/MasterCard", "TRUE"],
        [3, "محافظ إلكترونية", "📱", "Apple Pay / Google Pay", "TRUE"],
        [4, "دفع آمن", "🔒", "طرق الدفع والتوصيل تختلف حسب المتجر والمنصة.", "TRUE"]
    ]:
        ws.append(row)

    ws = wb.create_sheet("منصات")
    ws.append(["id", "name", "url", "showOnSite", "affiliateUrl", "notes"])
    for row in [
        [1, "Amazon", "#", "TRUE", "#", "Affiliate tagged link when available"],
        [2, "Temu", "#", "TRUE", "#", "منصة تسوق متنوعة"],
        [3, "Noon", "#", "TRUE", "#", "منصة تسوق محلية"]
    ]:
        ws.append(row)

    ws = wb.create_sheet("شرح التحكم")
    ws.append(["step", "instruction"])
    for row in [
        ["1", "فتح Excel"],
        ["2", "رفع الملف إلى Google Sheets"],
        ["3", "فتح Extensions"],
        ["4", "فتح Apps Script"],
        ["5", "وضع Code.gs"],
        ["6", "Deploy"],
        ["7", "Web App"],
        ["8", "نسخ رابط /exec"],
        ["9", "وضعه في config.js"],
        ["10", "تشغيل الموقع"]
    ]:
        ws.append(row)

    for sheet in wb.worksheets:
        for column_cells in sheet.columns:
            max_length = 0
            for cell in column_cells:
                try:
                    cell_len = len(str(cell.value))
                except Exception:
                    cell_len = 0
                if cell_len > max_length:
                    max_length = cell_len
            sheet.column_dimensions[column_cells[0].column_letter].width = max_length + 2

    out_path = OUT_DIR / "KOKO_SHOP_CONTROL_FINAL.xlsx"
    wb.save(out_path)
    return out_path


def write_guide_png() -> Path:
    width, height = 1400, 900
    img = Image.new("RGB", (width, height), color=(247, 249, 252))
    draw = ImageDraw.Draw(img)

    # background panel
    draw.rounded_rectangle((50, 50, 1350, 850), radius=30, fill=(255, 255, 255))
    draw.rounded_rectangle((100, 100, 1300, 780), radius=24, fill=(236, 244, 255))

    title = "KOKO SHOP - دليل التحكم"
    try:
        font_title = ImageFont.truetype("DejaVuSans-Bold.ttf", 44)
        font_step = ImageFont.truetype("DejaVuSans.ttf", 28)
        font_small = ImageFont.truetype("DejaVuSans.ttf", 22)
    except Exception:
        font_title = ImageFont.load_default()
        font_step = ImageFont.load_default()
        font_small = ImageFont.load_default()

    draw.text((120, 130), title, font=font_title, fill=(18, 33, 56))

    steps = [
        "1. فتح Excel",
        "2. رفعه إلى Google Sheets",
        "3. فتح Extensions",
        "4. فتح Apps Script",
        "5. وضع Code.gs",
        "6. Deploy",
        "7. Web App",
        "8. نسخ رابط /exec",
        "9. وضعه في config.js",
        "10. تشغيل الموقع",
    ]

    x0, y0 = 130, 210
    for idx, step in enumerate(steps, start=1):
        y = y0 + (idx - 1) * 52
        draw.rounded_rectangle((x0, y, x0 + 1040, y + 36), radius=10, fill=(255, 255, 255))
        draw.text((x0 + 16, y + 4), step, font=font_step, fill=(15, 94, 196))

    draw.text((120, 780), "ملاحظة: استخدم ملف Excel ثم ارفع إلى Google Sheets ثم اربط Web App في config.js", font=font_small, fill=(90, 110, 130))

    out_path = OUT_DIR / "KOKO_SHOP_CONTROL_GUIDE.png"
    img.save(out_path)
    return out_path


def write_zip_bundle() -> Path:
    bundle_name = "KOKO_SHOP_PREMIUM_GLOBAL_STYLE_FINAL.zip"
    bundle_path = OUT_DIR / bundle_name

    with zipfile.ZipFile(bundle_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for root, _, files in os.walk(ROOT):
            if ".git" in root or "dist" in root:
                continue
            for filename in files:
                if filename.endswith(".zip"):
                    continue
                full = Path(root) / filename
                arcname = full.relative_to(ROOT)
                zf.write(full, arcname)

    return bundle_path


def main() -> None:
    write_logo_jpg()
    write_excel()
    write_guide_png()
    zip_path = write_zip_bundle()
    print(f"Logo: {ROOT / 'assets' / 'koko-shop-logo.jpg'}")
    print(f"Excel: {OUT_DIR / 'KOKO_SHOP_CONTROL_FINAL.xlsx'}")
    print(f"Guide PNG: {OUT_DIR / 'KOKO_SHOP_CONTROL_GUIDE.png'}")
    print(f"ZIP: {zip_path}")


if __name__ == "__main__":
    main()
